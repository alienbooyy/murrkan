from sqlalchemy.orm import Session
from sqlalchemy import func, and_
from datetime import datetime, date
from typing import List, Optional
import json

from backend import models, schemas

# ===== TABLE CRUD =====
def get_tables(db: Session) -> List[models.Table]:
    return db.query(models.Table).all()

def get_table(db: Session, table_id: int) -> Optional[models.Table]:
    return db.query(models.Table).filter(models.Table.id == table_id).first()

def create_table(db: Session, table: schemas.TableCreate) -> models.Table:
    db_table = models.Table(**table.model_dump())
    db.add(db_table)
    db.commit()
    db.refresh(db_table)
    return db_table

def update_table(db: Session, table_id: int, table: schemas.TableUpdate) -> models.Table:
    db_table = get_table(db, table_id)
    if not db_table:
        raise ValueError("Table not found")
    
    update_data = table.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(db_table, key, value)
    
    db.commit()
    db.refresh(db_table)
    return db_table

def delete_table(db: Session, table_id: int):
    db_table = get_table(db, table_id)
    if db_table:
        db.delete(db_table)
        db.commit()

def merge_tables(db: Session, table_id: int, target_table_id: int):
    """Merge two tables - move orders from table_id to target_table_id"""
    source_table = get_table(db, table_id)
    target_table = get_table(db, target_table_id)
    
    if not source_table or not target_table:
        raise ValueError("Table not found")
    
    # Move all pending orders from source to target
    pending_orders = db.query(models.Order).filter(
        and_(
            models.Order.table_id == table_id,
            models.Order.status == models.OrderStatus.PENDING
        )
    ).all()
    
    for order in pending_orders:
        order.table_id = target_table_id
    
    # Mark source table as merged
    source_table.merged_with = target_table_id
    source_table.status = models.TableStatus.EMPTY
    
    # Mark target table as occupied
    target_table.status = models.TableStatus.OCCUPIED
    
    db.commit()
    return {"message": "Tables merged successfully"}

# ===== PRODUCT CRUD =====
def get_products(db: Session, active_only: bool = True) -> List[models.Product]:
    query = db.query(models.Product)
    if active_only:
        query = query.filter(models.Product.is_active == True)
    return query.all()

def get_product(db: Session, product_id: int) -> Optional[models.Product]:
    return db.query(models.Product).filter(models.Product.id == product_id).first()

def create_product(db: Session, product: schemas.ProductCreate) -> models.Product:
    db_product = models.Product(**product.model_dump())
    db.add(db_product)
    db.commit()
    db.refresh(db_product)
    return db_product

def update_product(db: Session, product_id: int, product: schemas.ProductUpdate) -> models.Product:
    db_product = get_product(db, product_id)
    if not db_product:
        raise ValueError("Product not found")
    
    update_data = product.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(db_product, key, value)
    
    db.commit()
    db.refresh(db_product)
    return db_product

def delete_product(db: Session, product_id: int):
    db_product = get_product(db, product_id)
    if db_product:
        db_product.is_active = False
        db.commit()

# ===== INGREDIENT CRUD =====
def get_ingredients(db: Session, active_only: bool = True) -> List[models.Ingredient]:
    query = db.query(models.Ingredient)
    if active_only:
        query = query.filter(models.Ingredient.is_active == True)
    return query.all()

def get_ingredient(db: Session, ingredient_id: int) -> Optional[models.Ingredient]:
    return db.query(models.Ingredient).filter(models.Ingredient.id == ingredient_id).first()

def create_ingredient(db: Session, ingredient: schemas.IngredientCreate) -> models.Ingredient:
    db_ingredient = models.Ingredient(**ingredient.model_dump())
    db.add(db_ingredient)
    db.commit()
    db.refresh(db_ingredient)
    return db_ingredient

def update_ingredient(db: Session, ingredient_id: int, ingredient: schemas.IngredientUpdate) -> models.Ingredient:
    db_ingredient = get_ingredient(db, ingredient_id)
    if not db_ingredient:
        raise ValueError("Ingredient not found")
    
    update_data = ingredient.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(db_ingredient, key, value)
    
    db.commit()
    db.refresh(db_ingredient)
    return db_ingredient

def delete_ingredient(db: Session, ingredient_id: int):
    db_ingredient = get_ingredient(db, ingredient_id)
    if db_ingredient:
        db_ingredient.is_active = False
        db.commit()

# ===== RECIPE CRUD =====
def get_product_recipe(db: Session, product_id: int) -> List[models.RecipeItem]:
    return db.query(models.RecipeItem).filter(models.RecipeItem.product_id == product_id).all()

def create_recipe_item(db: Session, recipe: schemas.RecipeItemCreate) -> models.RecipeItem:
    db_recipe = models.RecipeItem(**recipe.model_dump())
    db.add(db_recipe)
    db.commit()
    db.refresh(db_recipe)
    return db_recipe

def delete_recipe_item(db: Session, recipe_id: int):
    db_recipe = db.query(models.RecipeItem).filter(models.RecipeItem.id == recipe_id).first()
    if db_recipe:
        db.delete(db_recipe)
        db.commit()

# ===== ORDER CRUD =====
def get_orders(db: Session, status: str = None) -> List[models.Order]:
    query = db.query(models.Order)
    if status:
        query = query.filter(models.Order.status == status)
    return query.order_by(models.Order.created_at.desc()).all()

def get_order(db: Session, order_id: int) -> Optional[models.Order]:
    return db.query(models.Order).filter(models.Order.id == order_id).first()

def get_table_orders(db: Session, table_id: int) -> List[models.Order]:
    return db.query(models.Order).filter(
        and_(
            models.Order.table_id == table_id,
            models.Order.status == models.OrderStatus.PENDING
        )
    ).all()

def create_order(db: Session, order: schemas.OrderCreate) -> models.Order:
    db_order = models.Order(
        table_id=order.table_id,
        notes=order.notes
    )
    db.add(db_order)
    db.flush()
    
    # Add order items
    total = 0
    for item in order.items:
        product = get_product(db, item.product_id)
        if product:
            subtotal = product.price * item.quantity
            db_item = models.OrderItem(
                order_id=db_order.id,
                product_id=item.product_id,
                quantity=item.quantity,
                unit_price=product.price,
                subtotal=subtotal,
                notes=item.notes
            )
            db.add(db_item)
            total += subtotal
            
            # Update ingredient stock
            recipe_items = get_product_recipe(db, item.product_id)
            for recipe_item in recipe_items:
                ingredient = get_ingredient(db, recipe_item.ingredient_id)
                if ingredient:
                    ingredient.stock_quantity -= recipe_item.quantity * item.quantity
    
    db_order.total_amount = total
    
    # Update table status
    table = get_table(db, order.table_id)
    if table:
        table.status = models.TableStatus.OCCUPIED
    
    db.commit()
    db.refresh(db_order)
    return db_order

def update_order(db: Session, order_id: int, order: schemas.OrderUpdate) -> models.Order:
    db_order = get_order(db, order_id)
    if not db_order:
        raise ValueError("Order not found")
    
    update_data = order.model_dump(exclude_unset=True)
    for key, value in update_data.items():
        setattr(db_order, key, value)
    
    db.commit()
    db.refresh(db_order)
    return db_order

def add_order_item(db: Session, order_id: int, item: schemas.OrderItemCreate) -> models.OrderItem:
    order = get_order(db, order_id)
    if not order:
        raise ValueError("Order not found")
    
    product = get_product(db, item.product_id)
    if not product:
        raise ValueError("Product not found")
    
    subtotal = product.price * item.quantity
    db_item = models.OrderItem(
        order_id=order_id,
        product_id=item.product_id,
        quantity=item.quantity,
        unit_price=product.price,
        subtotal=subtotal,
        notes=item.notes
    )
    db.add(db_item)
    
    # Update order total
    order.total_amount += subtotal
    
    # Update ingredient stock
    recipe_items = get_product_recipe(db, item.product_id)
    for recipe_item in recipe_items:
        ingredient = get_ingredient(db, recipe_item.ingredient_id)
        if ingredient:
            ingredient.stock_quantity -= recipe_item.quantity * item.quantity
    
    db.commit()
    db.refresh(db_item)
    return db_item

def delete_order_item(db: Session, item_id: int):
    db_item = db.query(models.OrderItem).filter(models.OrderItem.id == item_id).first()
    if db_item:
        order = get_order(db, db_item.order_id)
        if order:
            order.total_amount -= db_item.subtotal
            
            # Restore ingredient stock
            product_id = db_item.product_id
            quantity = db_item.quantity
            recipe_items = get_product_recipe(db, product_id)
            for recipe_item in recipe_items:
                ingredient = get_ingredient(db, recipe_item.ingredient_id)
                if ingredient:
                    ingredient.stock_quantity += recipe_item.quantity * quantity
        
        db.delete(db_item)
        db.commit()

# ===== PAYMENT CRUD =====
def create_payment(db: Session, payment: schemas.PaymentCreate) -> models.Payment:
    db_payment = models.Payment(**payment.model_dump())
    db.add(db_payment)
    
    # Check if order is fully paid
    order = get_order(db, payment.order_id)
    if order:
        total_paid = db.query(func.sum(models.Payment.amount)).filter(
            models.Payment.order_id == payment.order_id
        ).scalar() or 0
        total_paid += payment.amount
        
        if total_paid >= order.total_amount:
            order.status = models.OrderStatus.COMPLETED
            order.completed_at = datetime.utcnow()
            order.payment_method = payment.payment_method
            
            # Update table status if no other pending orders
            pending_orders = db.query(models.Order).filter(
                and_(
                    models.Order.table_id == order.table_id,
                    models.Order.status == models.OrderStatus.PENDING,
                    models.Order.id != order.id
                )
            ).count()
            
            if pending_orders == 0:
                table = get_table(db, order.table_id)
                if table:
                    table.status = models.TableStatus.EMPTY
    
    db.commit()
    db.refresh(db_payment)
    return db_payment

def german_style_payment(db: Session, payment: schemas.GermanStylePaymentRequest):
    """German style payment - pay for specific items in an order"""
    order = get_order(db, payment.order_id)
    if not order:
        raise ValueError("Order not found")
    
    # Calculate total for selected items
    total = 0
    for item_id in payment.item_ids:
        item = db.query(models.OrderItem).filter(models.OrderItem.id == item_id).first()
        if item and item.order_id == payment.order_id:
            item.is_paid = True
            total += item.subtotal
    
    # Create payment
    db_payment = models.Payment(
        order_id=payment.order_id,
        amount=total,
        payment_method=payment.payment_method,
        notes="German style payment"
    )
    db.add(db_payment)
    
    # Check if all items are paid
    unpaid_items = db.query(models.OrderItem).filter(
        and_(
            models.OrderItem.order_id == payment.order_id,
            models.OrderItem.is_paid == False
        )
    ).count()
    
    if unpaid_items == 0:
        order.status = models.OrderStatus.COMPLETED
        order.completed_at = datetime.utcnow()
        order.payment_method = models.PaymentMethod.GERMAN_STYLE
        
        # Update table status
        pending_orders = db.query(models.Order).filter(
            and_(
                models.Order.table_id == order.table_id,
                models.Order.status == models.OrderStatus.PENDING
            )
        ).count()
        
        if pending_orders == 0:
            table = get_table(db, order.table_id)
            if table:
                table.status = models.TableStatus.EMPTY
    
    db.commit()
    return {"message": "Payment processed", "amount": total}

def close_order(db: Session, order_id: int, payment_method: str):
    """Close an order with full payment"""
    order = get_order(db, order_id)
    if not order:
        raise ValueError("Order not found")
    
    # Check if already paid
    total_paid = db.query(func.sum(models.Payment.amount)).filter(
        models.Payment.order_id == order_id
    ).scalar() or 0
    
    remaining = order.total_amount - total_paid
    
    if remaining > 0:
        # Create final payment
        db_payment = models.Payment(
            order_id=order_id,
            amount=remaining,
            payment_method=payment_method
        )
        db.add(db_payment)
    
    # Close order
    order.status = models.OrderStatus.COMPLETED
    order.completed_at = datetime.utcnow()
    order.payment_method = payment_method
    
    # Update table status
    pending_orders = db.query(models.Order).filter(
        and_(
            models.Order.table_id == order.table_id,
            models.Order.status == models.OrderStatus.PENDING,
            models.Order.id != order_id
        )
    ).count()
    
    if pending_orders == 0:
        table = get_table(db, order.table_id)
        if table:
            table.status = models.TableStatus.EMPTY
    
    db.commit()
    return {"message": "Order closed", "total": order.total_amount}

# ===== REPORT CRUD =====
def get_daily_report(db: Session, date_str: str):
    """Get daily report for a specific date"""
    target_date = datetime.strptime(date_str, "%Y-%m-%d").date()
    start_datetime = datetime.combine(target_date, datetime.min.time())
    end_datetime = datetime.combine(target_date, datetime.max.time())
    
    # Get completed orders for the day
    orders = db.query(models.Order).filter(
        and_(
            models.Order.completed_at >= start_datetime,
            models.Order.completed_at <= end_datetime,
            models.Order.status == models.OrderStatus.COMPLETED
        )
    ).all()
    
    total_revenue = sum(order.total_amount for order in orders)
    total_orders = len(orders)
    
    # Get product sales
    product_sales = {}
    total_cost = 0
    
    for order in orders:
        for item in order.order_items:
            product_id = item.product_id
            if product_id not in product_sales:
                product_sales[product_id] = {
                    "product_id": product_id,
                    "product_name": item.product.name,
                    "quantity_sold": 0,
                    "total_revenue": 0,
                    "cost": 0
                }
            
            product_sales[product_id]["quantity_sold"] += item.quantity
            product_sales[product_id]["total_revenue"] += item.subtotal
            
            # Calculate cost from recipe
            recipe_items = get_product_recipe(db, product_id)
            item_cost = sum(
                recipe.ingredient.cost_per_unit * recipe.quantity * item.quantity
                for recipe in recipe_items
            )
            product_sales[product_id]["cost"] += item_cost
            total_cost += item_cost
    
    profit = total_revenue - total_cost
    
    return {
        "date": date_str,
        "total_revenue": total_revenue,
        "total_orders": total_orders,
        "total_cost": total_cost,
        "profit": profit,
        "product_sales": list(product_sales.values())
    }

def get_date_range_report(db: Session, start_date: str, end_date: str):
    """Get report for a date range"""
    start_dt = datetime.strptime(start_date, "%Y-%m-%d")
    end_dt = datetime.strptime(end_date, "%Y-%m-%d")
    
    daily_reports = []
    total_revenue = 0
    total_orders = 0
    total_cost = 0
    
    current_date = start_dt
    while current_date <= end_dt:
        date_str = current_date.strftime("%Y-%m-%d")
        daily_report = get_daily_report(db, date_str)
        daily_reports.append(daily_report)
        total_revenue += daily_report["total_revenue"]
        total_orders += daily_report["total_orders"]
        total_cost += daily_report["total_cost"]
        current_date += timedelta(days=1)
    
    from datetime import timedelta
    
    return {
        "start_date": start_date,
        "end_date": end_date,
        "total_revenue": total_revenue,
        "total_orders": total_orders,
        "total_cost": total_cost,
        "profit": total_revenue - total_cost,
        "daily_breakdown": daily_reports
    }

def get_top_products(db: Session, limit: int = 10, start_date: str = None, end_date: str = None):
    """Get top selling products"""
    query = db.query(
        models.OrderItem.product_id,
        func.sum(models.OrderItem.quantity).label("total_quantity"),
        func.sum(models.OrderItem.subtotal).label("total_revenue")
    ).join(models.Order)
    
    if start_date and end_date:
        start_dt = datetime.strptime(start_date, "%Y-%m-%d")
        end_dt = datetime.strptime(end_date, "%Y-%m-%d")
        query = query.filter(
            and_(
                models.Order.completed_at >= start_dt,
                models.Order.completed_at <= end_dt
            )
        )
    
    query = query.filter(models.Order.status == models.OrderStatus.COMPLETED)
    query = query.group_by(models.OrderItem.product_id)
    query = query.order_by(func.sum(models.OrderItem.quantity).desc())
    query = query.limit(limit)
    
    results = []
    for row in query.all():
        product = get_product(db, row.product_id)
        if product:
            results.append({
                "product_id": row.product_id,
                "product_name": product.name,
                "quantity_sold": row.total_quantity,
                "total_revenue": row.total_revenue
            })
    
    return results

def export_daily_report_excel(db: Session, date_str: str):
    """Export daily report to Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    
    report = get_daily_report(db, date_str)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Report"
    
    # Header
    ws['A1'] = "GÜNLÜK RAPOR"
    ws['A1'].font = Font(size=16, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:D1')
    
    ws['A2'] = f"Tarih: {date_str}"
    ws['A2'].font = Font(size=12, bold=True)
    
    # Summary
    row = 4
    ws[f'A{row}'] = "Toplam Ciro:"
    ws[f'B{row}'] = report['total_revenue']
    ws[f'B{row}'].number_format = '#,##0.00 ₺'
    row += 1
    
    ws[f'A{row}'] = "Toplam Sipariş:"
    ws[f'B{row}'] = report['total_orders']
    row += 1
    
    ws[f'A{row}'] = "Toplam Maliyet:"
    ws[f'B{row}'] = report['total_cost']
    ws[f'B{row}'].number_format = '#,##0.00 ₺'
    row += 1
    
    ws[f'A{row}'] = "Kar:"
    ws[f'B{row}'] = report['profit']
    ws[f'B{row}'].number_format = '#,##0.00 ₺'
    ws[f'B{row}'].font = Font(bold=True)
    row += 2
    
    # Product sales table
    ws[f'A{row}'] = "Ürün Satışları"
    ws[f'A{row}'].font = Font(size=14, bold=True)
    row += 1
    
    headers = ["Ürün Adı", "Miktar", "Toplam Gelir"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = Font(bold=True)
    row += 1
    
    for product in report['product_sales']:
        ws.cell(row=row, column=1, value=product['product_name'])
        ws.cell(row=row, column=2, value=product['quantity_sold'])
        cell = ws.cell(row=row, column=3, value=product['total_revenue'])
        cell.number_format = '#,##0.00 ₺'
        row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 15
    
    # Save file
    import os
    os.makedirs("reports", exist_ok=True)
    filepath = f"reports/daily_report_{date_str}.xlsx"
    wb.save(filepath)
    
    return filepath
